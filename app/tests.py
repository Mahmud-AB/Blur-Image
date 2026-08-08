import json
import tempfile
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from .models import AnnotationRow, ImageUpload

User = get_user_model()


def generate_test_image(name="sample.jpg", image_format="JPEG", color="blue"):
    buffer = BytesIO()
    Image.new("RGB", (32, 32), color=color).save(buffer, format=image_format)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type=f"image/{image_format.lower()}",
    )


def _owned_image(user, file, original_name="sample.jpg"):
    return ImageUpload.objects.create(user=user, image=file, original_name=original_name)


def _response_bytes(response):
    if getattr(response, "streaming", False):
        return b"".join(response.streaming_content)
    return response.content


class HomePageTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="secret")
        self.client.force_login(self.user)

    def test_home_page_renders(self):
        response = self.client.get(reverse("home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Upload images")

    def test_upload_images_saves_files(self):
        response = self.client.post(reverse("upload_images"), {"images": [generate_test_image()]})

        self.assertEqual(response.status_code, 201)
        image = ImageUpload.objects.get()
        self.assertEqual(image.user, self.user)
        self.assertTrue(image.original_image)
        self.assertTrue(image.thumbnail)
        self.assertFalse(image.is_edited)

    def test_upload_images_preserves_bytes_and_resolution(self):
        buffer = BytesIO()
        Image.new("RGB", (1920, 1080), color="red").save(buffer, format="JPEG", quality=92)
        original_bytes = buffer.getvalue()
        file = SimpleUploadedFile("large.jpg", original_bytes, content_type="image/jpeg")

        response = self.client.post(reverse("upload_images"), {"images": [file]})

        self.assertEqual(response.status_code, 201)
        payload = response.json()["images"][0]
        self.assertEqual(payload["width"], 1920)
        self.assertEqual(payload["height"], 1080)
        self.assertEqual(payload["file_size"], len(original_bytes))

        image = ImageUpload.objects.get()
        image.image.open("rb")
        saved_bytes = image.image.read()
        image.image.close()
        self.assertEqual(saved_bytes, original_bytes)

    def test_delete_image_removes_saved_record(self):
        image = _owned_image(self.user, generate_test_image())
        response = self.client.post(reverse("delete_image", args=[image.id]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(ImageUpload.objects.count(), 0)

    def test_edit_image_replaces_saved_file(self):
        image = _owned_image(self.user, generate_test_image())
        response = self.client.post(
            reverse("edit_image", args=[image.id]),
            data=json.dumps({"points": [{"x": 0, "y": 0}, {"x": 10, "y": 0}, {"x": 10, "y": 10}]}),
            content_type="application/json",
        )

        image.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(image.image.name.endswith(".jpg"))
        self.assertTrue(image.is_edited)
        self.assertTrue(image.thumbnail)

    def test_edit_image_preserves_resolution(self):
        buffer = BytesIO()
        Image.new("RGB", (1920, 1080), color="green").save(buffer, format="JPEG", quality=95)
        original_file = SimpleUploadedFile("large.jpg", buffer.getvalue(), content_type="image/jpeg")
        image = _owned_image(self.user, original_file, original_name="large.jpg")

        response = self.client.post(
            reverse("edit_image", args=[image.id]),
            data=json.dumps(
                {"points": [{"x": 100, "y": 100}, {"x": 400, "y": 100}, {"x": 400, "y": 400}]}
            ),
            content_type="application/json",
        )

        image.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        with Image.open(image.image) as edited_image:
            self.assertEqual(edited_image.size, (1920, 1080))
            self.assertEqual(edited_image.format, "JPEG")

    def test_serve_image_returns_bytes_for_linked_user(self):
        image = _owned_image(self.user, generate_test_image())
        response = self.client.get(reverse("serve_image", args=[image.id]))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(_response_bytes(response))
        self.assertIn("image/", response["Content-Type"])

    def test_serve_image_download_uses_original_filename(self):
        image = _owned_image(
            self.user,
            generate_test_image(name="2_fp_img_0165.jpg"),
            original_name="2_fp_img_0165.jpg",
        )
        response = self.client.get(reverse("serve_image", args=[image.id]), {"download": "1"})

        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("2_fp_img_0165.jpg", response["Content-Disposition"])

    def test_serve_image_thumb_returns_smaller_jpeg(self):
        response = self.client.post(reverse("upload_images"), {"images": [generate_test_image()]})
        image_id = response.json()["images"][0]["id"]

        thumb = self.client.get(reverse("serve_image", args=[image_id]), {"thumb": "1"})
        self.assertEqual(thumb.status_code, 200)
        self.assertIn("image/", thumb["Content-Type"])
        with Image.open(BytesIO(_response_bytes(thumb))) as thumb_image:
            self.assertLessEqual(max(thumb_image.size), 216)

    def test_image_payload_uses_serve_url(self):
        response = self.client.post(reverse("upload_images"), {"images": [generate_test_image()]})
        payload = response.json()["images"][0]
        self.assertEqual(payload["url"], reverse("serve_image", args=[payload["id"]]))

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_edit_image_saves_annotation_coordinates_by_category(self):
        image = _owned_image(self.user, generate_test_image(color="blue"))
        response = self.client.post(
            reverse("edit_image", args=[image.id]),
            data=json.dumps(
                {
                    "category": "signboard",
                    "shapes": [[{"x": 0, "y": 0}, {"x": 10, "y": 0}, {"x": 10, "y": 10}, {"x": 0, "y": 10}]],
                }
            ),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        annotation = AnnotationRow.objects.get(
            user=self.user, category="signboard", image_name="sample.jpg"
        )
        self.assertEqual(annotation.coordinate_text, "0,0;10,0;10,10;0,10")

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_edit_image_updates_existing_annotation_row(self):
        image = _owned_image(self.user, generate_test_image(color="blue"))
        url = reverse("edit_image", args=[image.id])
        self.client.post(
            url,
            data=json.dumps(
                {"category": "signboard", "shapes": [[{"x": 0, "y": 0}, {"x": 10, "y": 0}, {"x": 10, "y": 10}]]}
            ),
            content_type="application/json",
        )
        self.client.post(
            url,
            data=json.dumps(
                {"category": "signboard", "shapes": [[{"x": 5, "y": 5}, {"x": 15, "y": 5}, {"x": 15, "y": 15}]]}
            ),
            content_type="application/json",
        )

        rows = AnnotationRow.objects.filter(user=self.user, category="signboard", image_name="sample.jpg")
        self.assertEqual(rows.count(), 1)
        self.assertEqual(rows.first().coordinate_text, "5,5;15,5;15,15")

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_download_annotations_xlsx(self):
        AnnotationRow.objects.create(
            user=self.user,
            category="signboard",
            image_name="sample.jpg",
            coordinate_text="0,0;10,0;10,10;0,10",
        )
        AnnotationRow.objects.create(
            user=self.user,
            category="road",
            image_name="sample.jpg",
            coordinate_text="20,20;30,20;30,30;20,30|40,40;50,40;50,50",
        )

        response = self.client.get(reverse("download_annotations"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('attachment; filename="annotations.xlsx"', response["Content-Disposition"])

        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("signboard", workbook.sheetnames)
        self.assertIn("road", workbook.sheetnames)
        signboard_rows = list(workbook["signboard"].iter_rows(values_only=True))
        road_rows = list(workbook["road"].iter_rows(values_only=True))
        self.assertEqual(signboard_rows[0], ("Image Name", "Coordinates"))
        self.assertIn(("sample.jpg", "0,0;10,0;10,10;0,10"), signboard_rows)
        self.assertIn(("sample.jpg", "20,20;30,20;30,30;20,30"), road_rows)
        self.assertIn(("sample.jpg", "40,40;50,40;50,50"), road_rows)

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_download_annotations_xlsx_for_single_image(self):
        image = _owned_image(self.user, generate_test_image(name="sample.jpg"))
        AnnotationRow.objects.create(
            user=self.user,
            category="signboard",
            image_name="sample.jpg",
            coordinate_text="0,0;10,0;10,10;0,10",
        )
        AnnotationRow.objects.create(
            user=self.user,
            category="road",
            image_name="other.jpg",
            coordinate_text="20,20;30,20;30,30;20,30",
        )

        response = self.client.get(reverse("download_annotations"), {"image_id": image.id})
        self.assertEqual(response.status_code, 200)
        self.assertIn("sample_annotations.xlsx", response["Content-Disposition"])

        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["signboard"])
        signboard_rows = list(workbook["signboard"].iter_rows(values_only=True))
        self.assertEqual(signboard_rows[0], ("Image Name", "Coordinates"))
        self.assertEqual(signboard_rows[1], ("sample.jpg", "0,0;10,0;10,10;0,10"))
        self.assertEqual(len(signboard_rows), 2)

    def test_restore_image_replaces_edited_version_with_original(self):
        original_file = generate_test_image(color="blue")
        original_bytes = original_file.read()
        image = ImageUpload(original_name="sample.jpg", user=self.user)
        image.image.save("sample.jpg", ContentFile(original_bytes), save=False)
        image.save()
        image.original_image.save("original_sample.jpg", ContentFile(original_bytes), save=False)
        image.save(update_fields=["original_image"])

        self.client.post(
            reverse("edit_image", args=[image.id]),
            data=json.dumps({"points": [{"x": 0, "y": 0}, {"x": 31, "y": 0}, {"x": 31, "y": 31}]}),
            content_type="application/json",
        )
        response = self.client.post(reverse("restore_image", args=[image.id]))

        image.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(image.image.name.endswith(".jpg"))
        self.assertFalse(image.is_edited)
